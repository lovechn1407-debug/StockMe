import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

TEMPLATE_PATH = r"C:\Users\Love Chauhan\Desktop\StockMe\Hackathon PPT Layout TechTrix.pptx"
OUTPUT_PATH = r"C:\Users\Love Chauhan\Desktop\StockMe\FoodStock_Hackathon_TechTrix.pptx"
ASSETS_DIR = r"C:\Users\Love Chauhan\Desktop\StockMe\presentation_assets"

prs = pptx.Presentation(TEMPLATE_PATH)

NAVY = RGBColor(31, 73, 125)
DARK = RGBColor(15, 23, 42)
SLATE = RGBColor(71, 85, 105)
BLUE = RGBColor(37, 99, 235)
GREEN = RGBColor(16, 185, 129)

# -------------------------------------------------------------
# SLIDE 1: Hackathon - TechTrix (Title Slide)
# -------------------------------------------------------------
slide1 = prs.slides[0]

# Update TextBox 15 (Project & Team Details)
tb15 = None
for s in slide1.shapes:
    if s.has_text_frame and "Title -" in s.text:
        tb15 = s
        break

if tb15:
    tf = tb15.text_frame
    tf.word_wrap = True
    tf.clear()

    # Title line
    p0 = tf.paragraphs[0]
    p0.text = "FoodStock"
    p0.font.name = "Georgia"
    p0.font.size = Pt(40)
    p0.font.bold = True
    p0.font.color.rgb = NAVY
    p0.space_after = Pt(6)

    # Subtitle
    p_sub = tf.add_paragraph()
    p_sub.text = "Smart Commodity Stock Allocation & FIFO Management in iPDS Godowns"
    p_sub.font.name = "Segoe UI"
    p_sub.font.size = Pt(18)
    p_sub.font.bold = True
    p_sub.font.color.rgb = BLUE
    p_sub.space_after = Pt(18)

    # PS Category
    p_cat = tf.add_paragraph()
    p_cat.text = "PS Category: Software (GovTech / Smart Supply Chain Automation)"
    p_cat.font.name = "Segoe UI"
    p_cat.font.size = Pt(16)
    p_cat.font.bold = True
    p_cat.font.color.rgb = DARK
    p_cat.space_after = Pt(12)

    # Team Name
    p_team = tf.add_paragraph()
    p_team.text = "Team Name (Registered): [Insert Team Name]"
    p_team.font.name = "Segoe UI"
    p_team.font.size = Pt(16)
    p_team.font.bold = True
    p_team.font.color.rgb = NAVY
    p_team.space_after = Pt(14)

    # Team Members
    p_mem_head = tf.add_paragraph()
    p_mem_head.text = "Team Member Details:"
    p_mem_head.font.name = "Segoe UI"
    p_mem_head.font.size = Pt(15)
    p_mem_head.font.bold = True
    p_mem_head.font.color.rgb = DARK
    p_mem_head.space_after = Pt(4)

    members = [
        "• Member 1 (Team Lead / Full Stack): [Name] • [Contact/Email] • [College]",
        "• Member 2 (Backend & Architecture): [Name] • [Contact/Email] • [College]",
        "• Member 3 (UI/UX & Frontend): [Name] • [Contact/Email] • [College]",
        "• Member 4 (Domain & PDS Research): [Name] • [Contact/Email] • [College]",
    ]
    for m in members:
        pm = tf.add_paragraph()
        pm.text = m
        pm.font.name = "Segoe UI"
        pm.font.size = Pt(13)
        pm.font.color.rgb = SLATE
        pm.space_after = Pt(3)

    # Position text box on the left half
    tb15.left = Inches(1.0)
    tb15.top = Inches(2.9)
    tb15.width = Inches(9.8)
    tb15.height = Inches(7.5)

# Add Hero Illustration on the right
hero_img_path = os.path.join(ASSETS_DIR, "hero_godown.jpg")
if os.path.exists(hero_img_path):
    slide1.shapes.add_picture(
        hero_img_path,
        left=Inches(11.2),
        top=Inches(2.9),
        width=Inches(7.8),
        height=Inches(6.8),
    )

# -------------------------------------------------------------
# SLIDE 2: IDEA TITLE & PROPOSED SOLUTION
# -------------------------------------------------------------
slide2 = prs.slides[1]

# Update group title
for s in slide2.shapes:
    if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
        for child in s.shapes:
            if child.has_text_frame and "IDEA TITLE" in child.text:
                child.text_frame.text = "IDEA: FOODSTOCK - AUTOMATED FIFO & MULTI-BATCH SLICING"

# Clear placeholder textbox and insert graphic
for s in list(slide2.shapes):
    if s.has_text_frame and "Proposed Solution" in s.text:
        # Hide original placeholder text
        s.text_frame.clear()

slide2_img = os.path.join(ASSETS_DIR, "slide2_fifo_diagram.png")
if os.path.exists(slide2_img):
    slide2.shapes.add_picture(
        slide2_img,
        left=Inches(1.0),
        top=Inches(3.4),
        width=Inches(18.0),
        height=Inches(7.2),
    )

# -------------------------------------------------------------
# SLIDE 3: TECHNICAL APPROACH
# -------------------------------------------------------------
slide3 = prs.slides[2]

# Clear placeholder textbox
for s in list(slide3.shapes):
    if s.has_text_frame and "Technologies to be used" in s.text:
        s.text_frame.clear()

slide3_img = os.path.join(ASSETS_DIR, "slide3_architecture.png")
if os.path.exists(slide3_img):
    slide3.shapes.add_picture(
        slide3_img,
        left=Inches(1.0),
        top=Inches(3.4),
        width=Inches(18.0),
        height=Inches(7.2),
    )

# -------------------------------------------------------------
# SLIDE 4: FEASIBILITY AND VIABILITY
# -------------------------------------------------------------
slide4 = prs.slides[3]

# Clear placeholder textbox
for s in list(slide4.shapes):
    if s.has_text_frame and "Analysis of the feasibility" in s.text:
        s.text_frame.clear()

slide4_img = os.path.join(ASSETS_DIR, "slide4_feasibility_matrix.png")
if os.path.exists(slide4_img):
    slide4.shapes.add_picture(
        slide4_img,
        left=Inches(1.0),
        top=Inches(3.4),
        width=Inches(18.0),
        height=Inches(7.2),
    )

# -------------------------------------------------------------
# SLIDE 5: IMPACT AND BENEFITS
# -------------------------------------------------------------
slide5 = prs.slides[4]

# Clear placeholder textbox
for s in list(slide5.shapes):
    if s.has_text_frame and "Potential impact on the target" in s.text:
        s.text_frame.clear()

slide5_img = os.path.join(ASSETS_DIR, "slide5_impact_metrics.png")
if os.path.exists(slide5_img):
    slide5.shapes.add_picture(
        slide5_img,
        left=Inches(1.0),
        top=Inches(3.4),
        width=Inches(18.0),
        height=Inches(7.2),
    )

# -------------------------------------------------------------
# SLIDE 6: RESEARCH AND REFERENCES
# -------------------------------------------------------------
slide6 = prs.slides[5]

# Clear placeholder textbox
for s in list(slide6.shapes):
    if s.has_text_frame and "Details / Links of the reference" in s.text:
        s.text_frame.clear()

slide6_img = os.path.join(ASSETS_DIR, "slide6_research_framework.png")
if os.path.exists(slide6_img):
    slide6.shapes.add_picture(
        slide6_img,
        left=Inches(1.0),
        top=Inches(3.4),
        width=Inches(18.0),
        height=Inches(7.2),
    )

# Save presentation
prs.save(OUTPUT_PATH)
print(f"Successfully generated populated presentation at: {OUTPUT_PATH}")
